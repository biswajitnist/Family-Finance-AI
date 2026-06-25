import csv
import hashlib
import io
import re
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

import openpyxl
from fastapi import HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.account import Account
from app.models.audit import StatementSet
from app.models.category import Category
from app.models.document import Document
from app.models.transaction import Transaction
from app.services.ai_service import classify_transaction
from app.services.duplicate_detection import find_duplicate_transaction_id
from app.services.investment_extractor import (
    extract_holdings,
    looks_like_portfolio,
    store_holdings,
)
from app.services.merchant_learning import apply_merchant_history, normalize_merchant
from app.services.ocr_service import extract_image, extract_pdf
from app.services.rule_engine import apply_rules_to_transaction
from app.services.statement_audit import (
    create_statement_set,
    record_document_audit,
)

ALLOWED_EXTENSIONS = {".pdf", ".csv", ".xlsx", ".xls", ".jpg", ".jpeg", ".png", ".txt"}
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y")
MOBILE_SCREENSHOT_DATE = re.compile(
    r"^\s*(?P<date>\d{2}\.\d{2}\.\d{4}|Yesterday)\s*[,|]?\s*$",
    re.IGNORECASE,
)
MOBILE_SCREENSHOT_AMOUNT = re.compile(
    r"(?P<amount>[+-]?(?:\d{1,3}(?:[.\s]\d{3})+|\d+)[,.]\d{2})"
    r"\s*[,|]?\s*$"
)
AMOUNT_LIKE = re.compile(r"(?<!\d)\d{1,3}(?:[.\s]\d{3})*[,.]\d{2}(?!\d)")
DATE_LIKE = re.compile(r"\b(\d{4}-\d{2}-\d{2}|\d{2}[./]\d{2}[./]\d{4})\b")
MOBILE_SCREENSHOT_IGNORED_LINES = (
    "transactions",
    "standing orders",
    "details",
    "discover contract overview",
    "pending transactions",
    "upcoming transactions",
    "in the next",
)
COMMERZBANK_BOOKING_DATE = re.compile(
    r"^\s*Buchungsdatum:\s*(?P<date>\d{2}\.\d{2}\.\d{4})\s*$"
)
COMMERZBANK_TRANSACTION = re.compile(
    r"^\s*(?P<vendor>.+?)\s+(?P<value_date>\d{2}\.\d{2})"
    r"\s+(?P<amount>\d{1,3}(?:\.\d{3})*,\d{2})(?P<debit>-)?\s*$"
)
CREDIT_CARD_TABLE_ROW = re.compile(
    r"^\s*(?P<date>\d{2}\.\d{2}\.\d{4})\s+"
    r"(?P<card>(?:\*{3,}|x{3,})\s*\d{3,4}|\d{3,4})\s+"
    r"(?P<description>.+?)\s+"
    r"(?P<amount>[+-]?\d{1,3}(?:[.\s]\d{3})*[,.]\d{2})\s*(?:€|EUR)?\s*$",
    re.IGNORECASE,
)
CREDIT_CARD_VERTICAL_AMOUNT = re.compile(
    r"^(?P<amount>[+-]?\d{1,3}(?:[.\s]\d{3})*[,.]\d{2})\s*(?:€|EUR)?$"
)
CREDIT_CARD_MASKED_CARD = re.compile(r"^(?:\*{3,}|x{3,})\s*\d{3,4}$", re.IGNORECASE)
COMMERZBANK_IGNORED_PREFIXES = (
    "VSA000",
    "Kontoauszug vom",
    "Auszug-Nr.",
    "IBAN:",
    "BIC :",
    "Beratungscenter",
    "Kaiserstr.",
    "60278 Frankfurt",
    "Ihr Ansprechpartner",
    "Team Privatkunden",
    "Telefonnummer",
    "Kontowährung",
    "zu Ihren Lasten",
    "Angaben zu den Umsätzen",
    "Folgeseite",
    "Alter Kontostand",
    "Neuer Kontostand",
    "Guthaben sind",
    "gungsfähig.",
    "Nähere Informationen",
    "entnommen werden.",
    "Der angegebene Kontostand",
    "Dies bedeutet",
    "stehenden Guthaben",
    "Somit können",
    "möglicherweise Zinsen",
    "eingeräumten oder",
)


def _safe_name(filename: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name).strip("._")
    return stem[:180] or "document"


async def store_document(
    db: Session,
    file: UploadFile,
    document_type: str,
    source_account_id: int | None,
) -> Document:
    filename = _safe_name(file.filename or "document")
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(415, f"Unsupported file type: {extension or 'unknown'}")
    if source_account_id and db.get(Account, source_account_id) is None:
        raise HTTPException(404, "Source account not found")
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"File exceeds {settings.max_upload_mb} MB limit")
    checksum = hashlib.sha256(content).hexdigest()
    stored_name = f"{uuid4().hex}_{filename}"
    path = settings.uploads_dir / stored_name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    document = Document(
        user_id=1,
        file_name=filename,
        file_path=str(path),
        file_type=extension.removeprefix("."),
        document_type=document_type,
        source_account_id=source_account_id,
        status="uploaded",
        validation_status="not_reviewed",
        checksum=checksum,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def _tabular_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        return [
            {
                (key or "").strip().lower(): str(value or "").strip()
                for key, value in row.items()
            }
            for row in csv.DictReader(io.StringIO(text))
        ]
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(item or "").strip().lower() for item in values[0]]
    return [
        {headers[index]: str(value or "").strip() for index, value in enumerate(row)}
        for row in values[1:]
    ]


def _mobile_screenshot_description(lines: list[str]) -> str:
    useful = []
    for line in lines:
        cleaned = line.strip(" \t,|")
        lowered = cleaned.casefold()
        if (
            not cleaned
            or any(prefix in lowered for prefix in MOBILE_SCREENSHOT_IGNORED_LINES)
            or re.fullmatch(r"[\W_]+", cleaned)
            or re.fullmatch(r"\d{1,2}:\d{2}.*", cleaned)
            or re.fullmatch(r"[A-Za-z]+\s+\d{4}", cleaned)
        ):
            continue
        useful.append(cleaned)
    return " ".join(useful).strip()


def _rows_from_mobile_screenshot_text(
    text: str, reference_date: date | None = None
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    pending: list[tuple[int, str]] = []
    raw_lines = list(enumerate(text.splitlines(), start=1))
    ordered_lines: list[tuple[int, str]] = []
    index = 0
    while index < len(raw_lines):
        current = raw_lines[index]
        if (
            MOBILE_SCREENSHOT_DATE.match(current[1])
            and index + 1 < len(raw_lines)
            and MOBILE_SCREENSHOT_AMOUNT.search(raw_lines[index + 1][1])
        ):
            ordered_lines.append(raw_lines[index + 1])
            ordered_lines.append(current)
            index += 2
            continue
        ordered_lines.append(current)
        index += 1

    for line_number, line in ordered_lines:
        if line.strip().casefold() == "transactions":
            pending = []
            continue
        if re.fullmatch(r"\s*[A-Za-z]+\s+\d{4}\s*", line):
            pending = []
            continue
        date_match = MOBILE_SCREENSHOT_DATE.match(line)
        if not date_match:
            pending.append((line_number, line))
            continue

        raw_date = date_match.group("date")
        if raw_date.casefold() == "yesterday":
            transaction_date = (
                reference_date - timedelta(days=1) if reference_date else None
            )
        else:
            transaction_date = _parse_date(raw_date)

        amount_index = None
        amount_match = None
        for index in range(len(pending) - 1, -1, -1):
            candidate = MOBILE_SCREENSHOT_AMOUNT.search(pending[index][1])
            if candidate:
                amount_index = index
                amount_match = candidate
                break

        if transaction_date and amount_index is not None and amount_match:
            amount_line_number, amount_line = pending[amount_index]
            previous_amount_index = next(
                (
                    index
                    for index in range(amount_index - 1, -1, -1)
                    if MOBILE_SCREENSHOT_AMOUNT.search(pending[index][1])
                ),
                -1,
            )
            description_lines = [
                item[1]
                for item in pending[previous_amount_index + 1 : amount_index]
            ]
            amount_prefix = amount_line[: amount_match.start()].strip()
            if len(re.sub(r"\W", "", amount_prefix)) > 2:
                description_lines.append(amount_prefix)
            description = _mobile_screenshot_description(description_lines)
            if description:
                amount = amount_match.group("amount").replace(" ", "")
                rows.append(
                    {
                        "date": transaction_date.isoformat(),
                        "vendor": description[:180],
                        "description": description,
                        "amount": amount,
                        "currency": "EUR",
                        "type": "debit" if amount.startswith("-") else "credit",
                        "_source_line_number": str(amount_line_number),
                        "_source_text": f"{description} {amount} {raw_date}",
                    }
                )
        pending = []
    return rows


def _rows_from_statement_text(
    text: str, reference_date: date | None = None
) -> list[dict[str, str]]:
    if "Kontoauszug vom" in text and "Angaben zu den Umsätzen" in text:
        rows = _rows_from_commerzbank_text(text)
        if rows:
            return rows
    table_rows = _rows_from_credit_card_table_text(text)
    if table_rows:
        return table_rows
    vertical_table_rows = _rows_from_vertical_credit_card_table_text(text)
    if vertical_table_rows:
        return vertical_table_rows
    rows: list[dict[str, str]] = []
    pattern = re.compile(
        r"^\s*(?P<date>\d{4}-\d{2}-\d{2}|\d{2}[./]\d{2}[./]\d{4})"
        r"\s+(?P<description>.+?)\s+"
        r"(?P<amount>-?\d[\d.,]*)\s*(?P<currency>[A-Z]{3})?\s*$"
    )
    for line_number, line in enumerate(text.splitlines(), start=1):
        match = pattern.match(line)
        if not match:
            continue
        description = match.group("description").strip()
        amount = match.group("amount")
        rows.append(
            {
                "date": match.group("date"),
                "vendor": description[:180],
                "description": description,
                "amount": amount,
                "currency": match.group("currency") or "EUR",
                "type": "debit" if amount.startswith("-") else "credit",
                "_source_line_number": str(line_number),
                "_source_text": line.strip(),
            }
        )
    return rows or _rows_from_mobile_screenshot_text(text, reference_date)


def _no_transaction_reason(
    text: str,
    document: Document,
    used_ocr: bool,
    use_local_ai: bool,
) -> str:
    clean_lines = [line.strip() for line in text.splitlines() if line.strip()]
    amount_like_count = sum(1 for line in clean_lines if AMOUNT_LIKE.search(line))
    date_like_count = sum(1 for line in clean_lines if DATE_LIKE.search(line))
    reasons = []
    if not text.strip():
        reasons.append("no readable text was extracted")
    elif not amount_like_count:
        reasons.append("no amount-like values were detected")
    elif not date_like_count:
        reasons.append(
            "amount-like values were found, but matching dates were not detected"
        )
    else:
        reasons.append(
            f"{amount_like_count} amount-like line(s) were found, but the layout "
            "did not match known statement patterns"
        )
    if not document.source_account_id:
        reasons.append(
            "no source account is selected, so rows cannot become review transactions"
        )
    if used_ocr:
        reasons.append(
            "OCR was used; low image quality can split dates, merchants, and amounts"
        )
    if not use_local_ai:
        reasons.append("Ollama fallback was disabled")
    return "No transactions extracted because " + "; ".join(reasons) + "."


def _reference_date_from_filename(filename: str) -> date | None:
    match = re.search(r"(?P<date>\d{4}-\d{2}-\d{2})", filename)
    return _parse_date(match.group("date")) if match else None


def _normalize_amount(value: str) -> Decimal:
    normalized = value.strip().replace(" ", "")
    if "," in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    return abs(Decimal(normalized)).quantize(Decimal("0.01"))


def _commerzbank_description(lines: list[str]) -> str:
    useful = []
    for line in lines:
        cleaned = line.strip()
        if not cleaned or cleaned.startswith(COMMERZBANK_IGNORED_PREFIXES):
            continue
        useful.append(cleaned)
    return " ".join(useful)[:2000]


def _commerzbank_reference(lines: list[str]) -> str | None:
    for index, line in enumerate(lines):
        if "End-to-End-Ref.:" not in line:
            continue
        value = line.split("End-to-End-Ref.:", 1)[1].strip()
        if value:
            return value[:120]
        if index + 1 < len(lines):
            next_line = lines[index + 1].strip()
            if next_line and ":" not in next_line:
                return next_line[:120]
    return None


def _rows_from_commerzbank_text(text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    current_booking_date: date | None = None
    pending: dict[str, object] | None = None

    def flush() -> None:
        nonlocal pending
        if pending is None:
            return
        detail_lines = pending.pop("detail_lines")
        assert isinstance(detail_lines, list)
        pending["description"] = _commerzbank_description(detail_lines)
        pending["reference_number"] = _commerzbank_reference(detail_lines)
        rows.append({key: str(value or "") for key, value in pending.items()})
        pending = None

    for line_number, line in enumerate(text.splitlines(), start=1):
        if line.strip().startswith("Folgeseite"):
            flush()
            continue
        booking_match = COMMERZBANK_BOOKING_DATE.match(line)
        if booking_match:
            flush()
            current_booking_date = _parse_date(booking_match.group("date"))
            continue

        transaction_match = COMMERZBANK_TRANSACTION.match(line)
        if transaction_match and current_booking_date:
            flush()
            value_day, value_month = map(
                int, transaction_match.group("value_date").split(".")
            )
            value_year = current_booking_date.year
            if value_month == 12 and current_booking_date.month == 1:
                value_year -= 1
            try:
                value_date = date(value_year, value_month, value_day)
            except ValueError:
                value_date = current_booking_date
            pending = {
                "date": current_booking_date.isoformat(),
                "booking_date": value_date.isoformat(),
                "vendor": transaction_match.group("vendor").strip()[:180],
                "amount": str(_normalize_amount(transaction_match.group("amount"))),
                "currency": "EUR",
                "type": ("debit" if transaction_match.group("debit") else "credit"),
                "confidence": "0.9800",
                "_source_line_number": str(line_number),
                "_source_text": line.strip(),
                "detail_lines": [],
            }
            continue

        if pending is not None:
            detail_lines = pending["detail_lines"]
            assert isinstance(detail_lines, list)
            detail_lines.append(line)

    flush()
    return rows


def _rows_from_credit_card_table_text(text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        cleaned = re.sub(r"\s+", " ", line).strip()
        if not cleaned or cleaned.casefold().startswith(("datum ", "karte ")):
            continue
        match = CREDIT_CARD_TABLE_ROW.match(cleaned)
        if not match:
            continue
        amount = match.group("amount")
        description = match.group("description").strip()
        rows.append(
            {
                "date": _parse_date(match.group("date")).isoformat(),
                "vendor": description[:180],
                "description": description,
                "amount": amount,
                "currency": "EUR",
                "type": "debit" if amount.startswith("-") else "credit",
                "payment_method": match.group("card").replace(" ", ""),
                "_source_line_number": str(line_number),
                "_source_text": cleaned,
                "confidence": "0.9600",
            }
        )
    return rows


def _rows_from_vertical_credit_card_table_text(text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    lines = [
        (line_number, line.strip())
        for line_number, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]
    index = 0
    while index < len(lines):
        line_number, current = lines[index]
        try:
            transaction_date = _parse_date(current)
        except ValueError:
            index += 1
            continue

        if index + 3 >= len(lines):
            break
        card = lines[index + 1][1]
        if not CREDIT_CARD_MASKED_CARD.match(card):
            index += 1
            continue

        description = lines[index + 2][1].strip()
        amount_match = CREDIT_CARD_VERTICAL_AMOUNT.match(lines[index + 3][1])
        if not description or not amount_match:
            index += 1
            continue

        amount = amount_match.group("amount")
        if _normalize_amount(amount) != Decimal("0.00"):
            rows.append(
                {
                    "date": transaction_date.isoformat(),
                    "vendor": description[:180],
                    "description": description,
                    "amount": amount,
                    "currency": "EUR",
                    "type": "debit" if amount.startswith("-") else "credit",
                    "payment_method": card.replace(" ", ""),
                    "_source_line_number": str(line_number),
                    "_source_text": " ".join(
                        [current, card, description, lines[index + 3][1]]
                    ),
                    "confidence": "0.9400",
                }
            )
        index += 5
    return rows


def _parse_date(value: str) -> date:
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(value.strip(), date_format).date()
        except ValueError:
            continue
    raise ValueError(f"Unsupported date: {value}")


def _create_transactions_from_rows(
    db: Session, document: Document, rows: list[dict[str, str]]
) -> list[tuple[dict[str, str], Transaction]]:
    if not document.source_account_id:
        return []
    categories = {
        item.name.casefold(): item.id for item in db.scalars(select(Category)).all()
    }
    created = []
    for row in rows:
        date_value = row.get("date") or row.get("transaction_date")
        vendor = row.get("vendor") or row.get("payee") or row.get("merchant")
        amount_value = row.get("amount")
        if not date_value or not vendor or not amount_value:
            continue
        try:
            amount = _normalize_amount(amount_value)
            transaction_date = _parse_date(date_value)
        except (InvalidOperation, ValueError):
            continue
        transaction_type = (
            row.get("type") or row.get("transaction_type") or ""
        ).lower()
        if transaction_type not in {"debit", "credit"}:
            transaction_type = (
                "debit" if amount_value.strip().startswith("-") else "credit"
            )
        description = row.get("description", "")
        transaction = Transaction(
            user_id=1,
            account_id=document.source_account_id,
            document_id=document.id,
            transaction_date=transaction_date,
            booking_date=_parse_date(row["booking_date"])
            if row.get("booking_date")
            else None,
            vendor=vendor,
            description=description,
            original_description=description,
            amount=amount,
            currency=(row.get("currency") or "EUR").upper(),
            transaction_type=transaction_type,
            category_id=categories.get(row.get("category", "").casefold()),
            payment_method=row.get("payment_method") or None,
            reference_number=row.get("reference_number") or None,
            iban=row.get("iban") or None,
            is_validated=False,
            is_duplicate=False,
            classification_source="import",
            confidence=Decimal(row["confidence"]).quantize(Decimal("0.0001"))
            if row.get("confidence")
            else None,
        )
        transaction.is_duplicate = (
            find_duplicate_transaction_id(db, transaction) is not None
        )
        db.add(transaction)
        db.flush()
        categorized = transaction.category_id is not None
        if not categorized:
            categorized = bool(apply_rules_to_transaction(db, transaction))
        if not categorized:
            categorized = bool(apply_merchant_history(db, transaction))
        if not categorized and not row.get("confidence"):
            classify_transaction(db, transaction)
        created.append((row, transaction))
    return created


def process_document(
    db: Session,
    document: Document,
    force_ocr: bool = False,
    use_local_ai: bool = True,
    statement_set_id: int | None = None,
    audit_only: bool = False,
) -> dict:
    path = Path(document.file_path)
    original_status = document.status
    original_validation_status = document.validation_status
    if any(item.is_validated for item in document.transactions) and not audit_only:
        raise HTTPException(409, "Confirmed documents cannot be processed again")
    if not audit_only:
        for transaction in list(document.transactions):
            db.delete(transaction)
    statement_set = (
        db.get(StatementSet, statement_set_id) if statement_set_id else None
    )
    if statement_set is None:
        statement_set = create_statement_set(db, [document])
    document.status = "processing"
    document.processing_stage = "preparing"
    document.processing_progress = 5
    document.processing_message = "Preparing the document for local extraction."
    document.processing_error = None
    db.commit()
    try:
        suffix = path.suffix.lower()
        rows: list[dict[str, str]] = []
        page_count = None
        used_ocr = False
        ocr_confidence = None
        ocr_lines = None
        document.processing_stage = "extracting"
        document.processing_progress = 20
        document.processing_message = "Reading document text."
        db.commit()
        if suffix == ".pdf":
            document.processing_message = (
                "Running OCR on the PDF."
                if force_ocr
                else "Reading PDF text and using OCR when needed."
            )
            db.commit()
            result = extract_pdf(path, force_ocr=force_ocr)
            text = result.text
            page_count = result.page_count
            used_ocr = result.used_ocr
            ocr_confidence = result.confidence
            ocr_lines = result.lines
        elif suffix in {".jpg", ".jpeg", ".png"}:
            document.processing_stage = "ocr"
            document.processing_message = "Running Tesseract OCR on the image."
            db.commit()
            result = extract_image(path)
            text = result.text
            page_count = result.page_count
            used_ocr = True
            ocr_confidence = result.confidence
            ocr_lines = result.lines
        elif suffix in {".csv", ".xlsx", ".xls"}:
            rows = _tabular_rows(path)
            for line_number, row in enumerate(rows, start=1):
                row.setdefault("confidence", "1.0000")
                row["_source_line_number"] = str(line_number)
                row["_source_text"] = " | ".join(
                    value
                    for key, value in row.items()
                    if not key.startswith("_")
                )
            text = "\n".join(row["_source_text"] for row in rows)
            statement_set.expected_transaction_count = (
                statement_set.expected_transaction_count or 0
            ) + len(rows)
            db.commit()
        else:
            text = path.read_text(encoding="utf-8-sig")
        if document.document_type == "investment_statement" or looks_like_portfolio(
            text
        ):
            document.document_type = "investment_statement"
            document.processing_stage = "parsing"
            document.processing_progress = 65
            document.processing_message = "Finding portfolio holdings."
            db.commit()
            holdings = extract_holdings(text)
            processed_path = settings.processed_dir / f"{document.id}.txt"
            processed_path.parent.mkdir(parents=True, exist_ok=True)
            processed_path.write_text(text, encoding="utf-8")
            created_holdings = store_holdings(db, document, holdings)
            document.extracted_text = text
            document.page_count = page_count
            document.status = (
                "validated" if created_holdings else "extraction_completed"
            )
            document.validation_status = (
                "validated" if created_holdings else "not_applicable"
            )
            document.processing_stage = "completed"
            document.processing_progress = 100
            document.processing_message = (
                f"{created_holdings} portfolio holding(s) added to Investments."
                if created_holdings
                else "Portfolio detected, but no holding rows were found."
            )
            document.processed_at = datetime.now(UTC)
            db.commit()
            return {
                "document_id": document.id,
                "status": document.status,
                "extracted_characters": len(text),
                "transactions_created": 0,
                "investments_created": created_holdings,
                "record_type": "investment",
                "message": (
                    "Portfolio holdings processed locally"
                    + (" with OCR" if used_ocr else "")
                ),
                "used_ocr": used_ocr,
                "ocr_confidence": ocr_confidence,
                "statement_set_id": statement_set.id,
            }
        document.processing_stage = "parsing"
        document.processing_progress = 55
        document.processing_message = "Finding transaction rows in the extracted text."
        db.commit()
        used_local_ai = False
        if not rows:
            rows = _rows_from_statement_text(
                text, reference_date=_reference_date_from_filename(document.file_name)
            )
        if not rows and use_local_ai and document.source_account_id:
            from app.services.ai_service import extract_transactions

            document.processing_stage = "local_ai"
            document.processing_progress = 70
            document.processing_message = "Asking local Ollama to structure the rows."
            db.commit()
            rows = extract_transactions(text, document.document_type)
            used_local_ai = True
        no_transaction_reason = (
            _no_transaction_reason(text, document, used_ocr, use_local_ai)
            if not rows
            else ""
        )
        processed_path = settings.processed_dir / f"{document.id}.txt"
        processed_path.parent.mkdir(parents=True, exist_ok=True)
        processed_path.write_text(text, encoding="utf-8")
        document.processing_stage = "creating_review"
        document.processing_progress = 85
        document.processing_message = "Creating review rows. Nothing is live yet."
        db.commit()
        if audit_only:
            existing = list(document.transactions)
            transaction_pairs = []
            for row in rows:
                amount = _normalize_amount(row.get("amount", "0"))
                match = next(
                    (
                        item
                        for item in existing
                        if item.transaction_date
                        == _parse_date(
                            row.get("date")
                            or row.get("transaction_date")
                            or ""
                        )
                        and item.amount == amount
                        and normalize_merchant(item.vendor)
                        == normalize_merchant(
                            row.get("vendor")
                            or row.get("merchant")
                            or row.get("payee")
                            or ""
                        )
                    ),
                    None,
                )
                if match:
                    transaction_pairs.append((row, match))
        else:
            transaction_pairs = _create_transactions_from_rows(db, document, rows)
        created = len(transaction_pairs)
        record_document_audit(
            db,
            statement_set,
            document,
            rows,
            transaction_pairs,
            text,
            ocr_lines,
        )
        no_review_rows_reason = ""
        if rows and not created and not audit_only and not document.source_account_id:
            no_review_rows_reason = (
                f"Found {len(rows)} transaction row(s), but no source account is "
                "selected, so review transactions were not created."
            )
        document.extracted_text = text
        document.page_count = page_count
        document.status = (
            original_status
            if audit_only
            else "needs_review"
            if created
            else "extraction_completed"
        )
        document.validation_status = (
            original_validation_status
            if audit_only
            else "pending"
            if created
            else "not_applicable"
        )
        document.processing_stage = (
            "audit_review"
            if audit_only
            else "human_review"
            if created
            else "completed"
        )
        document.processing_progress = 100
        document.processing_message = (
            no_transaction_reason
            if not rows
            else "Re-audit completed. Existing validated transactions were not changed."
            if audit_only
            else f"{created} row(s) are ready for human review."
            if created
            else no_review_rows_reason
            or "Rows were found, but no review transactions were created."
        )
        document.processed_at = datetime.now(UTC)
        db.commit()
        return {
            "document_id": document.id,
            "status": document.status,
            "extracted_characters": len(text),
            "transactions_created": created,
            "investments_created": 0,
            "record_type": "transaction",
            "message": (
                no_transaction_reason
                if not created and not rows
                else no_review_rows_reason
                if no_review_rows_reason
                else
                "Processing completed locally"
                + (" with OCR" if used_ocr else "")
                + (" and Ollama" if used_local_ai else "")
            ),
            "used_ocr": used_ocr,
            "ocr_confidence": ocr_confidence,
            "statement_set_id": statement_set.id,
        }
    except Exception as exc:
        document.status = "failed"
        document.processing_stage = "failed"
        document.processing_message = "Processing stopped. Review the error and retry."
        document.processing_error = str(exc)
        db.commit()
        raise HTTPException(422, f"Document processing failed: {exc}") from exc


def delete_document_files(document: Document) -> None:
    Path(document.file_path).unlink(missing_ok=True)
    processed = settings.processed_dir / f"{document.id}.txt"
    processed.unlink(missing_ok=True)
